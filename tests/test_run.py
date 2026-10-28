from collections import defaultdict
from datetime import timedelta
from decimal import Decimal

from openpyxl import load_workbook

from conftest import block, results, row
from tracebook.cli import main
from tracebook.core.notebook import parse, resolve
from tracebook.core.values import fmt, to_json
from tracebook.packs.razorpay.demo import generate


def _expected():
    """Independent pure-Python funnel from the generator's raw records (no SQL, no import path)."""
    d = generate()
    orders, pays = d["orders"], d["payments"]
    by_order = defaultdict(list)
    for p in pays:
        if p["order_id"]:
            by_order[p["order_id"]].append(p)
    start = min(o["created_at"] for o in orders).replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    settle_by = start + timedelta(days=29)        # May: 31 days − 2 (generator rule)
    paid = [o for o in orders if o["status"] == "paid"]
    exp = {
        "created": len(orders),
        "attempted": sum(o["attempts"] > 0 for o in orders),
        "paid": len(paid),
        "captured": sum(any(p["captured"] for p in by_order[o["id"]]) for o in orders),
        "settled": sum(any(p["captured"] and p["created_at"] < settle_by for p in by_order[o["id"]]) for o in orders),
        "lost": sum(Decimal(o["amount"]) for o in orders if o["status"] == "attempted"),
        "aov": sum(Decimal(o["amount"]) for o in paid) / len(paid),
    }
    m = defaultdict(lambda: {"payments": 0, "succeeded": 0, "lost_inr": Decimal(0)})
    for p in pays:
        if p["order_id"]:
            m[p["method"]]["payments"] += 1
            m[p["method"]]["succeeded"] += p["captured"]
    for o in orders:
        if o["status"] != "paid" and by_order[o["id"]]:
            last = max(by_order[o["id"]], key=lambda p: p["created_at"])
            m[last["method"]]["lost_inr"] += Decimal(o["amount"])
    exp["methods"] = dict(m)
    return exp


def test_demo_funnel_matches_independent_computation(demo):
    assert main(["check"]) == 0
    doc, exp = results(demo), _expected()
    f = row(doc, "funnel")
    assert {k: f[k] for k in ("created", "attempted", "paid", "captured", "settled")} == \
           {k: exp[k] for k in ("created", "attempted", "paid", "captured", "settled")}
    assert abs(f["conv_paid"] - exp["paid"] / exp["created"]) < 1e-12
    assert Decimal(row(doc, "lost")["inr"]) == exp["lost"]
    mb = block(doc, "methods")
    names = [c["name"] for c in mb["columns"]]
    got = {r[0]: dict(zip(names, r)) for r in mb["rows"]}
    for method, e in exp["methods"].items():
        assert got[method]["payments"] == e["payments"]
        assert got[method]["succeeded"] == e["succeeded"]
        assert Decimal(got[method]["lost_inr"]) == e["lost_inr"]
    card, w = exp["methods"]["card"], row(doc, "whatif")
    extra = round(card["payments"] * (0.75 - card["succeeded"] / card["payments"]))
    assert w["extra_payments"] == extra
    assert abs(w["uplift_inr"] - extra * float(exp["aov"])) < 1e-6
    weekly = block(doc, "weekly")
    assert sum(r[1] for r in weekly["rows"]) == exp["created"] and sum(r[2] for r in weekly["rows"]) == exp["paid"]
    md = block(doc, "md_1")["text"]
    assert f'data-ref="funnel.conv_paid">{fmt(f["conv_paid"], "pct")}<' in md


def test_sheet_model_source_untouched(demo):
    src = demo.root / "models" / "whatif.xlsx"
    before = src.read_bytes()
    assert main(["check"]) == 0
    assert src.read_bytes() == before
    copy = load_workbook(demo.nb().work_dir / "whatif.xlsx")
    assert copy["Inputs"]["B2"].value > 0                         # inputs were written to the copy only


def test_sheet_input_that_is_a_formula_is_rejected(demo):
    wb = load_workbook(demo.root / "models" / "whatif.xlsx")
    wb["Inputs"]["B2"] = "=1+1"
    wb.save(demo.root / "models" / "whatif.xlsx")
    assert main(["check"]) == 1
    doc = results(demo)
    assert block(doc, "whatif")["status"] == "error"
    assert any(d["code"] == "SHEET_INPUT_IS_FORMULA" for d in doc["diagnostics"])
    assert "#REF" in block(doc, [b["id"] for b in doc["blocks"] if b["kind"] == "md"][-1])["text"]


def _edit(p, old, new):
    t = p.notebook.read_text()
    assert old in t
    p.notebook.write_text(t.replace(old, new, 1))


def test_literal_number_in_prose_is_flagged(demo):
    _edit(demo, "# Checkout funnel", "# Checkout funnel\n\nConversion was 64.0% last month.")
    assert main(["check"]) == 1
    d = [d for d in results(demo)["diagnostics"] if d["code"] == "LITERAL_NUMBER_IN_PROSE"]
    assert len(d) == 1 and d[0]["line"] == 7 and "'64.0%'" in d[0]["message"]


def test_python_block_must_return_dataframe_with_line(demo):
    _edit(demo, 'out["conversion"] = out["paid"] / out["orders"]\nreturn out', 'return 42')
    assert main(["check"]) == 1
    doc = results(demo)
    assert block(doc, "weekly")["status"] == "error"
    assert all(b.get("error") for b in doc["blocks"] if b["kind"] == "chart" and b.get("data") == "weekly")
    assert any(d["code"] == "PY_NOT_DATAFRAME" for d in doc["diagnostics"])
    _edit(demo, "return 42", "return 1/0")
    main(["check"])
    err = next(d for d in results(demo)["diagnostics"] if d["code"] == "PY_ERROR")
    lines = demo.notebook.read_text().split("\n")
    assert lines[err["line"] - 1].strip() == "return 1/0"


def test_removed_block_leaves_no_stale_table(demo):
    assert main(["check"]) == 0
    _edit(demo, "```sql aov ", "```sql aov_renamed ")
    assert main(["check"]) == 1                                   # whatif still reads aov → REF_UNKNOWN
    doc = results(demo)
    assert any(d["code"] == "REF_UNKNOWN" and "aov" in d["message"] for d in doc["diagnostics"])
    con = demo.connect(read_only=True)
    tables = {t for (t,) in con.execute("SELECT table_name FROM duckdb_tables()").fetchall()}
    assert "aov" not in tables and "whatif" not in tables


def test_skipped_dependents_and_sql_error(demo):
    _edit(demo, "FROM orders WHERE status = 'attempted' AND currency = 'INR'",
          "FROM orders WHERE no_such_column = 1")
    assert main(["check"]) == 1
    doc = results(demo)
    assert block(doc, "lost")["status"] == "error"
    assert block(doc, "kpi_1")["items"][3]["value"] == "#REF"


def test_notebook_graph_errors():
    nb = parse("```sql a\nSELECT * FROM b\n```\n```sql b\nSELECT * FROM a\n```\n```sql c\nSELECT * FROM nope\n```\n"
               "```sql\nSELECT 1\n```\n```sql c\nSELECT 2\n```\n```bash\n```sql x\n```\n")
    order = resolve(nb, {"orders"})
    codes = sorted(d.code for d in nb.diags)
    assert codes == ["CYCLE", "CYCLE", "DUPLICATE_ID", "MISSING_ID", "REF_UNKNOWN"]
    assert nb.blocks[-1].kind == "md" and "```sql x" in nb.blocks[-1].body
    assert "x" not in {b.id for b in nb.blocks}                   # fenced code inside prose is not a block
    assert [b.id for b in order if b.kind == "sql"][:1] == ["c"]


def test_unclosed_fence():
    nb = parse("# t\n```sql a\nSELECT 1\n")
    assert [d.code for d in nb.diags] == ["UNCLOSED_FENCE"]


def test_values_exact():
    assert to_json(2**60) == str(2**60) and to_json(2**52) == 2**52
    assert to_json(Decimal("0.10")) == "0.10"
    assert fmt(182200, "inr") == "₹1,82,200" and fmt(1240, "inr") == "₹1,240" and fmt(-999, "inr") == "-₹999"
    assert fmt(Decimal("246000"), "inr_short") == "₹2.46L" and fmt(0.6395, "pct") == "64.0%"
