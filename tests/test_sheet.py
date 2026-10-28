import json

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from conftest import block, results
from tracebook.cli import main


def _model(path, **cells):
    wb = Workbook()
    ws = wb.active
    ws.title = "M"
    for a, v in cells.items():
        ws[a] = v
    wb.save(path)
    return path


def test_inspect_reports_values_and_errors(tmp_path, capsys):
    f = _model(tmp_path / "m.xlsx", A1=4, A2=0, B1="=A1*2", B2="=A1/A2", B3="=+A1")
    before = f.read_bytes()
    assert main(["--json", "sheet", str(f)]) == 1
    d = json.loads(capsys.readouterr().out)
    got = {c["a"]: (c["f"], c["v"]) for c in d["cells"]}
    assert got["B1"] == ("=A1*2", 8) and got["B2"] == ("=A1/A2", "#DIV/0!") and got["B3"] == ("=+A1", 4)
    assert [e["code"] for e in d["errors"]] == ["SHEET_FORMULA_ERROR"] and "M!B2" in d["errors"][0]["message"]
    assert d["sheets"] == [{"name": "M", "dims": "A1:B3", "formulas": 3}]
    assert f.read_bytes() == before                                # source never written


def test_inspect_range_and_limit(tmp_path, capsys):
    f = _model(tmp_path / "m.xlsx", A1=1, B1="=A1+1")
    assert main(["--json", "sheet", str(f), "--range", "M!A1:Z100000", "-n", "3"]) == 0
    d = json.loads(capsys.readouterr().out)
    assert [c["a"] for c in d["cells"]] == ["A1", "B1", "C1"] and d["truncated"]
    assert main(["sheet", str(f), "--range", "X!A1"]) == 1
    assert "BAD_SPEC" in capsys.readouterr().err


def test_diff_flags_removed_formula(tmp_path, capsys):
    old = _model(tmp_path / "old.xlsx", A1=1, B1="=A1+1", C1="=A1*3")
    new = _model(tmp_path / "new.xlsx", A1=1, B1=2, D1="=SUM(A1:C1)")
    wb = load_workbook(new)
    wb.create_sheet("Extra")
    wb.save(new)
    assert main(["--json", "sheet", str(new), "--diff", str(old)]) == 0
    d = json.loads(capsys.readouterr().out)
    assert [(c["cell"], c["old"], c["new"]) for c in d["changes"]] == \
           [("M!B1", "=A1+1", 2), ("M!C1", "=A1*3", None), ("M!D1", None, "=SUM(A1:C1)")]
    assert [w.split()[:2] for w in d["warnings"]] == [["FORMULA_REMOVED", "M!B1"], ["FORMULA_REMOVED", "M!C1"]]
    assert d["sheets_added"] == ["Extra"] and d["sheets_removed"] == []
    assert main(["--json", "sheet", str(old), "--diff", str(old)]) == 0
    assert json.loads(capsys.readouterr().out)["changes"] == []


def test_array_formula_input_is_rejected(demo):
    wb = load_workbook(demo.root / "models" / "whatif.xlsx")
    wb["Inputs"]["B2"] = ArrayFormula("B2", "=SUM(C2:C3)")
    wb.save(demo.root / "models" / "whatif.xlsx")
    assert main(["check"]) == 1
    doc = results(demo)
    assert block(doc, "whatif")["status"] == "error"
    assert any(d["code"] == "SHEET_INPUT_IS_FORMULA" for d in doc["diagnostics"])
