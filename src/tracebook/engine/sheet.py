"""`sheet` blocks: an .xlsx model (pi-sheets workflow) recalculated with formualizer on a copy.

The source model is never modified: inputs are written into .tracebook/models/<id>.xlsx, that copy is
recalculated, and outputs + the `show` range are read back.
"""
from __future__ import annotations

import re
import tempfile
import warnings
from datetime import datetime
from decimal import Decimal
from itertools import islice
from pathlib import Path

import formualizer as fz
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

# Same table as pi-sheets recalc_xlsx.py (formualizer 0.5.x error kinds → Excel codes).
KIND_TO_CODE = {"Div": "#DIV/0!", "Name": "#NAME?", "Ref": "#REF!", "Na": "#N/A", "Null": "#NULL!", "Num": "#NUM!",
                "Value": "#VALUE!", "Spill": "#SPILL!", "Calc": "#CALC!", "Error": "#ERROR!", "Cancelled": "#ERROR!",
                "Circ": "#REF!", "Nimpl": "#NAME?"}
_CELL = re.compile(r"^(?:'((?:[^']|'')+)'|([^!]+))!\$?([A-Z]{1,3})\$?(\d+)$")


class SheetError(Exception):
    def __init__(self, code: str, msg: str):
        super().__init__(msg)
        self.code = code


def split_ref(ref: str) -> tuple[str, str]:
    m = _CELL.match(str(ref).strip())
    if not m:
        raise SheetError("BAD_SPEC", f"cell reference {ref!r} must look like Sheet!B2 or 'My Sheet'!B2")
    return (m[1].replace("''", "'") if m[1] else m[2]), f"{m[3]}{m[4]}"


def _strip_unary_plus(f: str) -> str:
    # pi-sheets workaround: formualizer 0.5.8 returns #VALUE! for `=+<string cell>`; Excel passes through.
    return "=" + f[2:] if f.startswith("=+") and len(f) > 2 and f[2] not in "+-=" else f


def _excel_value(v):
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    if isinstance(v, datetime) and v.tzinfo:
        return v.replace(tzinfo=None)          # xlsx has no time zones; value is already in project TZ
    return v


def excel_display(v, nf: str | None) -> str:
    """Display string for a cell using the common Excel number formats (percent, grouping, fixed decimals)."""
    if v is None:
        return ""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    f = (nf or "").split(";")[0]
    places = len(f.split(".")[1].rstrip("%")) if "." in f else 0
    if "%" in f:
        return f"{v * 100:.{places}f}%"
    if "0" in f:
        return f"{v:,.{places}f}" if "," in f else f"{v:.{places}f}"
    return str(int(v)) if float(v).is_integer() else f"{v:.6g}"


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(path)


def _content(v):
    return getattr(v, "text", v)               # array/data-table formulas are objects; .text is the formula


def _range(wb, ref: str):
    sheet, _, rng = str(ref).rpartition("!")
    sheet = sheet.strip("'").replace("''", "'")
    if sheet not in wb.sheetnames:
        raise SheetError("BAD_SPEC", f"range {ref}: sheet {sheet!r} not in model (has {', '.join(wb.sheetnames)})")
    try:
        b = range_boundaries(rng)
    except (ValueError, TypeError):
        b = (None,)
    if None in b:
        raise SheetError("BAD_SPEC", f"range {ref!r} must look like Sheet!A1:D10")
    return wb[sheet], *b


def _recalc(wb, copy: Path):
    """Save wb (with the `=+` workaround) to copy and evaluate it.
    → (formulas {(sheet, row, col): original formula}, value(sheet, row, col), errors [(code, msg)])"""
    formulas = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                f = _content(c.value)
                if isinstance(f, str) and f.startswith("="):
                    formulas[(ws.title, c.row, c.column)] = f
                    if isinstance(c.value, str):
                        c.value = _strip_unary_plus(f)
    copy.parent.mkdir(parents=True, exist_ok=True)
    wb.save(copy)
    try:
        fwb = fz.Workbook.load_path(str(copy))
        fwb.evaluate_all()
    except Exception as e:
        raise SheetError("SHEET_FORMULA_ERROR", f"formualizer could not evaluate the model: {e}")

    def value(sheet, r, c):
        if (sheet, r, c) in formulas:
            v = fwb.evaluate_cell(sheet, r, c)
            if isinstance(v, dict) and v.get("type") == "Error":
                return {"error": KIND_TO_CODE.get(v.get("kind", ""), "#ERROR!"), "kind": v.get("kind")}
            return v
        return _content(wb[sheet].cell(r, c).value)

    errors = []
    for (sheet, r, c), f in formulas.items():
        v = value(sheet, r, c)
        if isinstance(v, dict) and "error" in v:
            code = "SHEET_UNSUPPORTED_FN" if v["kind"] == "Nimpl" else "SHEET_FORMULA_ERROR"
            errors.append((code, f"{sheet}!{wb[sheet].cell(r, c).coordinate} {f} → {v['error']}"))
    return formulas, value, errors


def _cell(ws, r, c, formulas, value) -> dict:
    cell = ws.cell(r, c)
    v = value(ws.title, r, c)
    v = v["error"] if isinstance(v, dict) else v
    nf = cell.number_format if cell.number_format != "General" else None
    return {"a": cell.coordinate, "v": v, "f": formulas.get((ws.title, r, c)), "fmt": nf, "d": excel_display(v, nf)}


def run(spec: dict, root: Path, work: Path, bid: str, inputs: dict[str, object]) -> dict:
    """inputs: cell ref → resolved python value. → {outputs: {name: value}, grid, errors}"""
    if not isinstance(spec, dict) or not spec.get("file"):
        raise SheetError("BAD_SPEC", "sheet block needs `file:` (and usually inputs/outputs/show)")
    src = (root / spec["file"]).resolve()
    if not src.is_relative_to(root) or not src.exists():
        raise SheetError("SHEET_FILE_MISSING", f"model file {spec['file']!r} not found in the project")
    wb = _load(src)
    for ref, val in inputs.items():
        sheet, coord = split_ref(ref)
        if sheet not in wb.sheetnames:
            raise SheetError("BAD_SPEC", f"input {ref}: sheet {sheet!r} not in model (has {', '.join(wb.sheetnames)})")
        cell = wb[sheet][coord]
        if isinstance(f := _content(cell.value), str) and f.startswith("="):
            raise SheetError("SHEET_INPUT_IS_FORMULA", f"input {ref} holds formula {f!r}; inputs must be constant cells")
        cell.value = _excel_value(val)
    formulas, value, errors = _recalc(wb, work / f"{bid}.xlsx")
    outputs = {}
    for ref, name in (spec.get("outputs") or {}).items():
        sheet, coord = split_ref(ref)
        if sheet not in wb.sheetnames:
            raise SheetError("BAD_SPEC", f"output {ref}: sheet {sheet!r} not in model")
        cell = wb[sheet][coord]
        v = value(sheet, cell.row, cell.column)
        outputs[str(name)] = None if isinstance(v, dict) else v
    grid = None
    if spec.get("show"):
        ws, c1, r1, c2, r2 = _range(wb, spec["show"])
        if (r2 - r1 + 1) * (c2 - c1 + 1) > 2000:
            raise SheetError("BAD_SPEC", f"show range {spec['show']} is larger than 2,000 cells")
        grid = {"sheet": ws.title, "range": spec["show"].rpartition("!")[2], "rows":
                [[_cell(ws, r, c, formulas, value) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]}
    return {"outputs": outputs, "grid": grid, "errors": errors}


def inspect(path: Path, ref: str | None = None, limit: int = 500) -> dict:
    """Standalone model check: recalc a temp copy (source untouched) → sheets, non-empty cells
    (formula, computed value), formula errors. `ref` limits cells to one range."""
    wb = _load(path)
    with tempfile.TemporaryDirectory() as tmp:
        formulas, value, errors = _recalc(wb, Path(tmp) / "model.xlsx")
        if ref:
            ws, c1, r1, c2, r2 = _range(wb, ref)
            total, keys = (r2 - r1 + 1) * (c2 - c1 + 1), ((ws, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1))
        else:
            keys = [(ws, c.row, c.column) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                    if c.value is not None]
            total = len(keys)
        cells = [{"sheet": ws.title, **_cell(ws, r, c, formulas, value)} for ws, r, c in islice(keys, limit)]
    sheets = [{"name": ws.title, "dims": ws.dimensions, "formulas": sum(k[0] == ws.title for k in formulas)}
              for ws in wb.worksheets]
    return {"sheets": sheets, "cells": cells, "truncated": total > limit,
            "errors": [{"code": c, "message": m} for c, m in errors]}


def diff(old: Path, new: Path) -> dict:
    """Cell-content diff (formula text or constant) between two workbooks; warns when a formula is lost."""
    a, b = _load(old), _load(new)
    contents = lambda ws: {(c.row, c.column): _content(c.value) for row in ws.iter_rows() for c in row
                           if c.value is not None}
    isf = lambda v: isinstance(v, str) and v.startswith("=")
    changes, warns = [], []
    for name in [s for s in b.sheetnames if s in a.sheetnames]:
        x, y = contents(a[name]), contents(b[name])
        for k in sorted(x.keys() | y.keys()):
            if x.get(k) == y.get(k):
                continue
            coord = b[name].cell(*k).coordinate
            changes.append({"cell": f"{name}!{coord}", "old": x.get(k), "new": y.get(k)})
            if isf(x.get(k)) and not isf(y.get(k)):
                warns.append(f"FORMULA_REMOVED {name}!{coord} {x[k]} replaced by "
                             f"{'nothing' if y.get(k) is None else repr(y[k])}")
    removed = [s for s in a.sheetnames if s not in b.sheetnames]
    warns += [f"SHEET_REMOVED {s} ({sum(isf(v) for v in contents(a[s]).values())} formula(s))" for s in removed]
    return {"sheets_added": [s for s in b.sheetnames if s not in a.sheetnames], "sheets_removed": removed,
            "changes": changes, "warnings": warns}
